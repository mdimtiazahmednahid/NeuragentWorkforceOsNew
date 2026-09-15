import os
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    os.system("pip install python-docx")
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Define styles
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

# Title
title = doc.add_heading('Case Study: Deploying Workforce OS to Shared Hosting', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Subtitle / Author
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Authored by: Md Imtiaz Ahmed Nahid\n')
run.bold = True
run.font.size = Pt(12)
run_date = subtitle.add_run('Date: September 2026\n')
run_date.font.color.rgb = RGBColor(100, 100, 100)

doc.add_heading('1. Executive Summary', level=1)
p = doc.add_paragraph(
    "Deploying a modern, monolithic Laravel 11 application (Workforce OS) onto a heavily restricted "
    "shared hosting environment (InfinityFree) presented significant DevOps, database, and asset-compilation challenges. "
    "This case study outlines the critical bottlenecks encountered and the engineered solutions implemented to "
    "achieve a seamless production deployment without SSH or terminal access."
)

doc.add_heading('2. Challenge: Database Key Length Restrictions', level=2)
p = doc.add_paragraph()
p.add_run("Problem:\n").bold = True
p.add_run("InfinityFree utilizes an older MariaDB architecture restricting index keys to 1000 bytes. Laravel’s default utf8mb4 encoding (255 characters) exceeded this limit during the creation of composite indexes on tables such as failed_jobs.\n")
p.add_run("Solution:\n").bold = True
p.add_run("Injected Schema::defaultStringLength(191); into the AppServiceProvider boot method and explicitly constrained UUID string lengths in the database migrations to adhere to legacy database limitations.")

doc.add_heading('3. Challenge: Security Restrictions on Symlinks', level=2)
p = doc.add_paragraph()
p.add_run("Problem:\n").bold = True
p.add_run("Shared hosting environments frequently disable the PHP symlink() function to prevent directory traversal attacks. This broke Laravel's core file storage mechanism, causing 500 Internal Server Errors when fetching user avatars.\n")
p.add_run("Solution:\n").bold = True
p.add_run("Engineered a bypass by modifying config/filesystems.php, mapping the 'public' disk root directly to the public_path('storage') directory, entirely eliminating the symlink dependency.")

doc.add_heading('4. Challenge: Mixed Content (HTTP vs HTTPS) & Mobile Rendering', level=2)
p = doc.add_paragraph()
p.add_run("Problem:\n").bold = True
p.add_run("The application rendered perfectly on desktop but collapsed into raw HTML on mobile devices (specifically iOS Safari). InfinityFree operates behind a proxy that terminates SSL, causing Laravel to generate HTTP links for CSS assets. Strict mobile browsers blocked these mixed-content assets.\n")
p.add_run("Solution:\n").bold = True
p.add_run("Forced HTTPS routing unconditionally via URL::forceScheme('https') in the application service provider and dynamically corrected the APP_URL environment variable to reflect the true secure origin.")

doc.add_heading('5. Challenge: Vite Development Server Conflict', level=2)
p = doc.add_paragraph()
p.add_run("Problem:\n").bold = True
p.add_run("A residual 'public/hot' file generated during local development was inadvertently migrated. This signaled Laravel Vite to bypass production assets and poll localhost:5173 for CSS, completely severing styling on non-local devices.\n")
p.add_run("Solution:\n").bold = True
p.add_run("Developed a custom PHP cleanup script (hot_fix.php) to seek and destroy the invisible 'hot' file directly on the production server, immediately restoring the compiled Vite CSS manifest.")

doc.add_heading('6. Challenge: Database Reconstruction without SSH', level=2)
p = doc.add_paragraph()
p.add_run("Problem:\n").bold = True
p.add_run("Extensive schema debugging required complete database wipes on production. Re-seeding complex relational data (Users, Roles, Tasks, Departments) without terminal access was impossible.\n")
p.add_run("Solution:\n").bold = True
p.add_run("Developed an automated suite of web-accessible PHP scripts (restore_all.php, taskadd.php) that parsed predefined arrays of critical business data and executed raw Eloquent insertions directly into the remote MySQL instance, successfully restoring operational parity in seconds.")

doc.add_heading('7. Conclusion', level=1)
p = doc.add_paragraph(
    "Successfully bridging the gap between local enterprise development and restricted shared hosting required deep "
    "framework knowledge and unorthodox problem-solving. By dynamically patching service providers, intercepting routing, "
    "and automating remote database seeding via HTTP requests, Workforce OS was fully stabilized and optimized for production."
)

doc.save('/Users/MacBookPro/Downloads/workforceOs/Case_Study_Workforce_OS.docx')
print("Successfully generated Case_Study_Workforce_OS.docx")
